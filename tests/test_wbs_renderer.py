from __future__ import annotations

import copy
import importlib.util
import json
import shutil
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
VALIDATOR = ROOT / "skills" / "work-wbs" / "scripts" / "validate_wbs.py"
RENDERER = ROOT / "skills" / "work-wbs" / "scripts" / "render_wbs.py"
TEMPLATE = ROOT / "skills" / "work-wbs" / "assets" / "lukuku-wbs-template.html"
STYLESHEET = ROOT / "skills" / "work-wbs" / "assets" / "lukuku-wbs-template.css"
FIXTURE = ROOT / "tests" / "fixtures" / "wbs" / "sample-wbs.json"


def load_validator():
    spec = importlib.util.spec_from_file_location("work_wbs_validator", VALIDATOR)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def manifest() -> dict:
    return json.loads(FIXTURE.read_text())


def load_renderer():
    spec = importlib.util.spec_from_file_location("work_wbs_renderer", RENDERER)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_sample_manifest_is_valid() -> None:
    validator = load_validator()
    validator.validate_manifest(manifest())


@pytest.mark.parametrize(
    ("mutate", "message"),
    [
        (lambda data: data["items"][1].update(display_id="TASK-001"), "duplicate display ID"),
        (lambda data: data["items"][0].update(dependencies=["DEV-2"]), "dependency cycle"),
        (
            lambda data: data["items"][0]["dates"]["baseline"].update(
                start="2026-09-23", end="2026-09-22"
            ),
            "ends before it starts",
        ),
        (
            lambda data: data["items"][1]["actual_effort"].update(source_ref=None),
            "actual effort requires a source",
        ),
        (lambda data: data["items"][0].update(result_evidence=[]), "completed item requires result evidence"),
        (lambda data: data["items"][2]["estimate"].update(max=1), "milestone effort must be zero"),
    ],
)
def test_invalid_manifest_is_rejected(mutate, message: str) -> None:
    validator = load_validator()
    data = copy.deepcopy(manifest())
    mutate(data)
    with pytest.raises(ValueError, match=message):
        validator.validate_manifest(data)


def test_summary_preserves_missing_zero_ranges_and_units() -> None:
    validator = load_validator()
    summary = validator.summarize_manifest(manifest())

    assert summary["task_count"] == 2
    assert summary["milestone_count"] == 1
    assert summary["estimated_effort"] == {"MH": {"min": 20, "max": 28}}
    assert summary["actual_effort"] == {"MH": 10}
    assert summary["missing_actual_count"] == 0
    assert summary["unscheduled_ids"] == ["TASK-003"]
    assert summary["unresolved_count"] == 1


def test_summary_keeps_incompatible_units_separate() -> None:
    validator = load_validator()
    data = copy.deepcopy(manifest())
    data["items"][1]["estimate"]["unit"] = "days"
    summary = validator.summarize_manifest(data)

    assert summary["estimated_effort"] == {
        "MH": {"min": 8, "max": 12},
        "days": {"min": 12, "max": 16},
    }


def test_renderer_creates_editable_workbook_without_losing_zero(tmp_path: Path) -> None:
    renderer = load_renderer()
    output = tmp_path / "verish-wbs.xlsx"
    renderer.render_workbook(manifest(), output)

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["기본정보", "WBS", "문서이력"]
    sheet = workbook["WBS"]
    assert sheet.freeze_panes == "A5"
    assert sheet.auto_filter.ref.startswith("A4:")
    assert sheet.page_setup.orientation == "landscape"
    rows = {sheet.cell(row, 1).value: row for row in range(5, sheet.max_row + 1)}
    assert set(rows) == {"TASK-001", "TASK-002", "MS-001", "TASK-003"}
    assert sheet.cell(rows["TASK-002"], 20).value == 0
    assert sheet.cell(rows["MS-001"], 20).value is None
    assert workbook["기본정보"]["B13"].value == 2
    assert workbook["문서이력"]["A4"].value == "v0.1"


def test_renderer_creates_safe_standard_html(tmp_path: Path) -> None:
    renderer = load_renderer()
    output = tmp_path / "verish-wbs.html"
    css = tmp_path / STYLESHEET.name
    shutil.copy2(STYLESHEET, css)
    renderer.render_html(manifest(), TEMPLATE, css, output)
    text = output.read_text()

    assert "LUKUKU-WBS-TEMPLATE-11" in text
    assert "{{" not in text
    assert "<script" not in text.lower()
    assert "https://" not in text
    positions = [
        text.index(f'data-wbs-section="{name}"')
        for name in (
            "version-history",
            "overall-schedule",
            "monthly-calendar",
            "weekly-calendar",
            "milestones",
            "work-detail",
            "acceptance-detail",
        )
    ]
    assert positions == sorted(positions)
    for display_id in ("TASK-001", "TASK-002", "TASK-003", "MS-001"):
        assert text.count(f'data-item-id="{display_id}"') >= 1


def test_renderer_cli_creates_all_artifacts(tmp_path: Path) -> None:
    output = tmp_path / "artifacts"
    result = subprocess.run(
        ["python3", str(RENDERER), str(FIXTURE), str(output)],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert (output / "verish-wbs.xlsx").is_file()
    assert (output / "verish-wbs.pdf").read_bytes().startswith(b"%PDF")
    assert json.loads((output / "wbs-manifest.json").read_text())["document"][
        "number"
    ] == "LUKUKU-WBS-VERISH-001"
