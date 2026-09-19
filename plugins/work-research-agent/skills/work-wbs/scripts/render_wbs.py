#!/usr/bin/env python3
from __future__ import annotations

import argparse
import calendar
import html
import importlib.util
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
_validator_spec = importlib.util.spec_from_file_location(
    "work_wbs_validate", SCRIPT_DIR / "validate_wbs.py"
)
assert _validator_spec is not None and _validator_spec.loader is not None
_validator = importlib.util.module_from_spec(_validator_spec)
_validator_spec.loader.exec_module(_validator)
load_manifest = _validator.load_manifest
validate_manifest = _validator.validate_manifest
summarize_manifest = _validator.summarize_manifest


DARK = "161C29"
MINT = "22DF88"
LIGHT_MINT = "EAFFF5"
LIGHT_GRAY = "F7F9FB"
STATUS_FILLS = {
    "예정": "F1F3F5",
    "진행 중": "DCFFF0",
    "확인 대기": "FFF4CB",
    "완료": "BDEFD8",
    "보류": "E4E7EB",
}
WBS_HEADERS = [
    "ID", "작업 키", "작업명", "구분", "단계", "업무 영역", "담당", "협업",
    "상태", "일정 기준", "기준 시작", "기준 종료", "현재 시작", "현재 종료",
    "실제 시작", "실제 종료", "예상 최소", "예상 최대", "예상 단위", "실제 공수",
    "실제 단위", "착수 조건", "선행 작업", "완료 기준", "요구사항", "결과 근거",
    "변경 사유", "다음 조치", "확정 범위", "출처 구분",
]


def _date_value(value: str | None) -> date | None:
    return date.fromisoformat(value) if value else None


def _period(item: dict[str, Any]) -> tuple[date | None, date | None]:
    category = "actual" if item["wbs_status"] == "완료" else "forecast"
    period = item["dates"][category]
    if not period.get("start") and not period.get("end"):
        period = item["dates"]["baseline"]
    return _date_value(period.get("start")), _date_value(period.get("end"))


def _format_estimate(estimate: dict[str, Any] | None) -> str:
    if estimate is None:
        return "미입력"
    minimum, maximum, unit = estimate["min"], estimate["max"], estimate["unit"]
    return f"{minimum:g} {unit}" if minimum == maximum else f"{minimum:g}-{maximum:g} {unit}"


def _format_effort(value: dict[str, Any] | None) -> str:
    if value is None:
        return "미입력"
    return f"{value['value']:g} {value['unit']}"


def _safe_slug(value: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9_-]+", "-", value).strip("-").lower()
    if not slug:
        raise ValueError("project slug has no safe filename characters")
    return slug


def render_workbook(manifest: dict[str, Any], output: Path) -> None:
    validate_manifest(manifest)
    if output.exists():
        raise ValueError(f"refusing to overwrite existing workbook: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    summary = summarize_manifest(manifest)
    workbook = Workbook()
    info = workbook.active
    info.title = "기본정보"
    wbs = workbook.create_sheet("WBS")
    history = workbook.create_sheet("문서이력")

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.4
        sheet.page_margins.bottom = 0.4

    info.merge_cells("A1:D1")
    info["A1"] = "프로젝트 WBS"
    info["A1"].font = Font(size=20, bold=True, color=DARK)
    fields = [
        ("프로젝트", manifest["document"]["project"]),
        ("고객 / 조직", manifest["document"]["customer"]),
        ("담당 PM", manifest["document"]["pm"]),
        ("문서 번호", manifest["document"]["number"]),
        ("문서 버전", manifest["document"]["version"]),
        ("표준 양식", manifest["template_version"]),
        ("현황 기준일", manifest["document"]["as_of_date"]),
        ("문서 상태", manifest["document"]["status"]),
        ("대상 범위", manifest["scope"]["title"]),
    ]
    for row, (label, value) in enumerate(fields, start=3):
        info.cell(row, 1, label).font = Font(bold=True)
        info.cell(row, 2, value)
    info["A13"] = "등록 작업 수"
    info["B13"] = summary["task_count"]
    info["A14"] = "마일스톤 수"
    info["B14"] = summary["milestone_count"]
    info["A15"] = "미확정 항목 수"
    info["B15"] = summary["unresolved_count"]
    info["A16"] = "일정 미정 ID"
    info["B16"] = ", ".join(summary["unscheduled_ids"]) or "없음"
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70
    info.print_area = "A1:D18"

    wbs.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(WBS_HEADERS))
    wbs.cell(1, 1, f"{manifest['document']['project']} · 작업별 일정·공수")
    wbs.cell(1, 1).font = Font(size=16, bold=True, color=DARK)
    wbs.cell(2, 1, f"현황 기준일 {manifest['document']['as_of_date']}")
    for column, header in enumerate(WBS_HEADERS, start=1):
        cell = wbs.cell(4, column, header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=DARK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row_number, item in enumerate(manifest["items"], start=5):
        estimate = item.get("estimate") or {}
        actual_effort = item.get("actual_effort") or {}
        dates = item["dates"]
        criteria = [entry["text"] for entry in item["completion_criteria"]]
        values = [
            item["display_id"], item["work_item_key"], item["title"], item["item_type"],
            item["phase"], item["workstream"], item.get("owner"), ", ".join(item.get("collaborators") or []),
            item["wbs_status"], item["schedule_basis"], dates["baseline"].get("start"), dates["baseline"].get("end"),
            dates["forecast"].get("start"), dates["forecast"].get("end"), dates["actual"].get("start"), dates["actual"].get("end"),
            estimate.get("min"), estimate.get("max"), estimate.get("unit"), actual_effort.get("value"), actual_effort.get("unit"),
            "\n".join(item.get("prerequisites") or []), ", ".join(item.get("dependencies") or []), "\n".join(criteria),
            ", ".join(item.get("requirement_refs") or []), ", ".join(item.get("result_evidence") or []),
            item.get("change_note"), item.get("next_action"), "예" if item["confirmed_scope"] else "아니오", item["origin"],
        ]
        for column, value in enumerate(values, start=1):
            cell = wbs.cell(row_number, column, value)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_number % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        wbs.cell(row_number, 9).fill = PatternFill("solid", fgColor=STATUS_FILLS[item["wbs_status"]])
    wbs.freeze_panes = "A5"
    wbs.auto_filter.ref = f"A4:{get_column_letter(len(WBS_HEADERS))}{wbs.max_row}"
    wbs.print_title_rows = "1:4"
    wbs.print_area = f"A1:{get_column_letter(len(WBS_HEADERS))}{wbs.max_row}"
    for column in range(1, len(WBS_HEADERS) + 1):
        wbs.column_dimensions[get_column_letter(column)].width = 14
    wbs.column_dimensions["C"].width = 28
    wbs.column_dimensions["X"].width = 34

    history.merge_cells("A1:G1")
    history["A1"] = "문서 버전 및 변경 이력"
    history["A1"].font = Font(size=16, bold=True, color=DARK)
    history_headers = ["버전", "변경일", "변경 위치·작업 ID", "주요 변경 내용·사유", "작성자", "확인·합의 기록", "근거"]
    for column, header in enumerate(history_headers, start=1):
        cell = history.cell(3, column, header)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=DARK)
    for row_number, change in enumerate(manifest["change_history"], start=4):
        values = [change["version"], change["date"], ", ".join(change["affected_ids"]), change["reason"], change["author"], change["confirmation"], change["source_ref"]]
        for column, value in enumerate(values, start=1):
            history.cell(row_number, column, value).alignment = Alignment(wrap_text=True, vertical="top")
    history.freeze_panes = "A4"
    history.auto_filter.ref = f"A3:G{history.max_row}"
    history.print_area = f"A1:G{history.max_row}"
    for column, width in enumerate((14, 16, 28, 45, 18, 28, 18), start=1):
        history.column_dimensions[get_column_letter(column)].width = width

    workbook.save(output)


def _page(manifest: dict[str, Any], title: str, body: str, number: int, total: int, *, section: str | None = None, cover: bool = False) -> str:
    section_attr = f' data-wbs-section="{section}"' if section else ""
    css = "page cover" if cover else "page"
    doc = manifest["document"]
    return (
        f'<section class="{css}"{section_attr}>'
        '<header class="header"><span class="brand">Lukuku.</span><span>프로젝트 WBS / 공유용 실행 계획</span></header>'
        f'<main class="content">{body}</main>'
        f'<footer class="footer"><span>{html.escape(doc["project"])} / 문서 {html.escape(doc["version"])} / 양식 {html.escape(manifest["template_version"])} / 현황 기준일 {html.escape(doc["as_of_date"])}</span><span>{number} / {total}</span></footer>'
        '</section>'
    )


def _table(headers: list[str], rows: list[list[str]], classes: str = "") -> str:
    head = "".join(f"<th>{html.escape(value)}</th>" for value in headers)
    body = "".join("<tr>" + "".join(f"<td>{value}</td>" for value in row) + "</tr>" for row in rows)
    return f'<table class="{classes}"><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>'


def _month_range(start: str, end: str) -> list[tuple[int, int]]:
    year, month = map(int, start.split("-"))
    end_year, end_month = map(int, end.split("-"))
    values = []
    while (year, month) <= (end_year, end_month):
        values.append((year, month))
        month += 1
        if month == 13:
            year += 1
            month = 1
    return values


def _calendar_body(manifest: dict[str, Any], year: int, month: int) -> str:
    events: dict[date, list[dict[str, Any]]] = defaultdict(list)
    for item in manifest["items"]:
        start, end = _period(item)
        for value in {start, end}:
            if value and value.year == year and value.month == month:
                events[value].append(item)
    cells = [f'<div class="calendar-head">{label}</div>' for label in ("월", "화", "수", "목", "금", "토", "일")]
    for week in calendar.Calendar(firstweekday=0).monthdatescalendar(year, month):
        for day in week:
            muted = " muted" if day.month != month else ""
            entries = "".join(
                f'<span class="calendar-event" data-item-id="{html.escape(item["display_id"])}">{html.escape(item["display_id"])} {html.escape(item["title"])}</span>'
                for item in events.get(day, [])
            )
            cells.append(f'<div class="{muted.strip()}"><span class="calendar-day">{day.day}</span>{entries}</div>')
    return f'<h2>{year}년 {month}월</h2><div class="calendar">{"".join(cells)}</div>'


def build_html_content(manifest: dict[str, Any]) -> list[tuple[str | None, str, bool]]:
    doc = manifest["document"]
    summary = summarize_manifest(manifest)
    pages: list[tuple[str | None, str, bool]] = []
    cover = (
        '<div class="cover-kicker">PROJECT DELIVERY PLAN</div><div class="cover-title">프로젝트 WBS</div>'
        f'<h1>{html.escape(doc["project"])}</h1><div class="accent"></div>'
        f'<p>{html.escape(manifest["scope"]["title"])}</p>'
        '<div class="cover-grid">'
        f'<div class="card"><div class="card-label">고객 / 조직</div><div class="card-value">{html.escape(doc["customer"])}</div></div>'
        f'<div class="card"><div class="card-label">담당 PM</div><div class="card-value">{html.escape(doc["pm"])}</div></div>'
        f'<div class="card"><div class="card-label">문서 버전</div><div class="card-value">{html.escape(doc["version"])}</div></div>'
        f'<div class="card"><div class="card-label">현황 기준일</div><div class="card-value">{html.escape(doc["as_of_date"])}</div></div>'
        '</div><div class="notice">초안은 공유나 개발 승인을 의미하지 않습니다. 미확정 값은 임의로 채우지 않습니다.</div>'
    )
    pages.append((None, cover, True))

    contents = (
        '<h2>목차</h2><div class="contents-list">'
        '<div><span>01</span><strong>전체 일정</strong></div>'
        '<div><span>02</span><strong>월간 일정표</strong></div>'
        '<div><span>03</span><strong>주간 일정표</strong></div>'
        '<div><span>04</span><strong>주요 마일스톤</strong></div>'
        '<div><span>05</span><strong>작업별 일정·공수</strong></div>'
        '<div><span>06</span><strong>착수 조건·완료 기준</strong></div>'
        '</div><div class="notice">문서 버전 및 변경 이력은 본문 앞에서 확인할 수 있습니다.</div>'
    )
    pages.append(("contents", contents, False))

    history_rows = [[html.escape(str(c[k])) if k != "affected_ids" else html.escape(", ".join(c[k])) for k in ("version", "date", "affected_ids", "reason", "author", "confirmation", "source_ref")] for c in manifest["change_history"]]
    pages.append(("version-history", '<h2>문서 버전 및 변경 이력</h2>' + _table(["버전", "변경일", "변경 위치·작업 ID", "주요 변경 내용·사유", "작성자", "확인·합의 기록", "근거"], history_rows), False))

    effort = ", ".join(f"{value['min']:g}-{value['max']:g} {unit}" for unit, value in summary["estimated_effort"].items()) or "미입력"
    actual = ", ".join(f"{value:g} {unit}" for unit, value in summary["actual_effort"].items()) or "미입력"
    attention = ", ".join(summary["unscheduled_ids"]) or "없음"
    phases = []
    seen = set()
    for item in manifest["items"]:
        if item["phase"] not in seen:
            seen.add(item["phase"]); phases.append(item["phase"])
    start_year, start_month = map(int, manifest["rendering"]["start_month"].split("-"))
    gantt_start = date(start_year, start_month, 1)
    gantt_start -= timedelta(days=gantt_start.weekday())
    gantt_cells = ['<div class="gantt-corner">단계</div>']
    for week in range(16):
        week_start = gantt_start + timedelta(days=week * 7)
        gantt_cells.append(f'<div class="gantt-week">{week_start:%m/%d}</div>')
    for phase in phases:
        phase_periods = [
            _period(item) for item in manifest["items"]
            if item["phase"] == phase and all(_period(item))
        ]
        phase_start = min((start for start, _ in phase_periods), default=None)
        phase_end = max((end for _, end in phase_periods), default=None)
        start_value = phase_start.isoformat() if phase_start else ""
        end_value = phase_end.isoformat() if phase_end else ""
        gantt_cells.append(
            f'<div class="gantt-label" data-gantt-phase="{html.escape(phase, quote=True)}" '
            f'data-gantt-start="{start_value}" data-gantt-end="{end_value}">{html.escape(phase)}</div>'
        )
        for week in range(16):
            week_start = gantt_start + timedelta(days=week * 7)
            week_end = week_start + timedelta(days=6)
            active = phase_start is not None and phase_end is not None and phase_start <= week_end and phase_end >= week_start
            gantt_cells.append('<div class="gantt-bar"></div>' if active else '<div class="gantt-empty"></div>')
    overview = (
        '<h2>01 전체 일정</h2><div class="summary-grid">'
        f'<div class="card"><div class="card-label">현황 기준일</div><div class="card-value">{html.escape(doc["as_of_date"])}</div></div>'
        f'<div class="card"><div class="card-label">등록 예상 공수</div><div class="card-value">{html.escape(effort)}</div></div>'
        f'<div class="card"><div class="card-label">실제 누적 공수</div><div class="card-value">{html.escape(actual)}</div></div>'
        f'<div class="card"><div class="card-label">작업 완료</div><div class="card-value">{summary["status_counts"].get("완료", 0)}/{summary["task_count"]}</div></div>'
        f'</div><div class="notice warning">먼저 확인할 항목: {html.escape(attention)}</div><div class="gantt">{"".join(gantt_cells)}</div>'
        '<p class="muted small">단계 막대는 기간 범위를 보여주며 공수나 진행률을 뜻하지 않습니다.</p>'
    )
    pages.append(("overall-schedule", overview, False))

    for index, (year, month) in enumerate(_month_range(manifest["rendering"]["start_month"], manifest["rendering"]["end_month"])):
        pages.append(("monthly-calendar" if index == 0 else None, _calendar_body(manifest, year, month), False))

    as_of = date.fromisoformat(doc["as_of_date"])
    monday = as_of - timedelta(days=as_of.weekday())
    weekly_parts = ['<h2>03 주간 일정표</h2>']
    for offset, label in ((0, "기준일이 속한 주"), (7, "다음 주")):
        start = monday + timedelta(days=offset); end = start + timedelta(days=6)
        rows = []
        for item in manifest["items"]:
            item_start, item_end = _period(item)
            if item_start and item_end and item_start <= end and item_end >= start:
                rows.append([f'<span data-item-id="{html.escape(item["display_id"])}">{html.escape(item["display_id"])}</span>', html.escape(item["title"]), html.escape(item_start.isoformat()), html.escape(item_end.isoformat()), html.escape(item["wbs_status"])])
        weekly_parts.append(f'<h3>{start.isoformat()} - {end.isoformat()} | {label}</h3>' + _table(["ID", "작업", "시작", "종료", "상태"], rows or [["-", "해당 작업 없음", "-", "-", "-"]]))
    pages.append(("weekly-calendar", "".join(weekly_parts), False))

    milestones = [item for item in manifest["items"] if item["item_type"] == "milestone"]
    milestone_rows = []
    for item in milestones:
        baseline = item["dates"]["baseline"]; forecast = item["dates"]["forecast"]
        criterion = "; ".join(entry["text"] for entry in item["completion_criteria"])
        milestone_rows.append([f'<span data-item-id="{html.escape(item["display_id"])}">{html.escape(item["display_id"])}</span>', html.escape(item["title"]), html.escape(baseline.get("end") or "미정"), html.escape(forecast.get("end") or "미정"), html.escape(item["wbs_status"]), html.escape(criterion)])
    application_rows = [
        ["적용 범위", html.escape(" / ".join(manifest["scope"]["inclusions"]))],
        ["제외 범위", html.escape(" / ".join(manifest["scope"]["exclusions"]) or "없음")],
        ["미확정 범위", html.escape(" / ".join(manifest["scope"]["unresolved"]) or "없음")],
        ["공수 해석", "예상 범위와 실제 누적을 분리하며, 미입력은 0과 다릅니다."],
    ]
    pages.append(("milestones", '<h2>04 주요 마일스톤과 적용 기준</h2>' + _table(["ID", "마일스톤", "기준 종료일", "현재 / 실제 종료일", "상태", "완료 기준"], milestone_rows) + _table(["구분", "내용"], application_rows), False))

    chunk_size = 9
    for offset in range(0, len(manifest["items"]), chunk_size):
        rows = []
        for item in manifest["items"][offset:offset + chunk_size]:
            dates = item["dates"]
            rows.append([
                f'<span data-item-id="{html.escape(item["display_id"])}">{html.escape(item["display_id"])}</span>',
                html.escape(item["title"]), html.escape(item.get("owner") or "미정"),
                html.escape(f"{dates['baseline'].get('start') or '미정'} / {dates['baseline'].get('end') or '미정'}"),
                html.escape(f"{dates['forecast'].get('start') or '미정'} / {dates['forecast'].get('end') or '미정'}"),
                html.escape(f"{dates['actual'].get('start') or '미입력'} / {dates['actual'].get('end') or '미입력'}"),
                html.escape(_format_estimate(item.get("estimate"))), html.escape(_format_effort(item.get("actual_effort"))),
                f'<span class="status-{html.escape(item["wbs_status"].replace(" ", "-"))}">{html.escape(item["wbs_status"])}</span>',
            ])
        pages.append(("work-detail" if offset == 0 else None, '<h2>05 작업별 일정·공수</h2>' + _table(["ID", "작업", "담당", "기준 시작/종료", "현재 시작/종료", "실제 시작/종료", "예상", "실제", "진행"], rows, "compact"), False))

    for offset in range(0, len(manifest["items"]), 8):
        rows = []
        for item in manifest["items"][offset:offset + 8]:
            criteria = "; ".join(entry["text"] for entry in item["completion_criteria"])
            rows.append([
                f'<span data-item-id="{html.escape(item["display_id"])}">{html.escape(item["display_id"])}</span>',
                html.escape("; ".join(item.get("prerequisites") or []) or "선행 조건 없음"),
                html.escape(criteria), html.escape(", ".join(item.get("requirement_refs") or []) or "미입력"),
                html.escape(", ".join(item.get("result_evidence") or []) or "미입력"),
                html.escape(item.get("change_note") or "-") + (f'<br>{html.escape(item.get("next_action") or "")}' if item.get("next_action") else ""),
            ])
        pages.append(("acceptance-detail" if offset == 0 else None, '<h2>06 착수 조건·완료 기준</h2>' + _table(["작업", "선행 작업·착수 조건", "완료 기준", "관련 요구사항", "결과 근거", "변경 사유·다음 조치"], rows, "compact"), False))
    return pages


def render_html(manifest: dict[str, Any], template: Path, stylesheet: Path, output: Path) -> None:
    validate_manifest(manifest)
    if output.exists():
        raise ValueError(f"refusing to overwrite existing HTML: {output}")
    if not template.is_file() or not stylesheet.is_file():
        raise ValueError("WBS template assets are missing")
    pages = build_html_content(manifest)
    total = len(pages)
    content = "".join(_page(manifest, manifest["document"]["project"], body, index, total, section=section, cover=cover) for index, (section, body, cover) in enumerate(pages, 1))
    text = template.read_text(encoding="utf-8")
    text = text.replace("{{TITLE}}", html.escape(f"WBS · {manifest['document']['project']}", quote=True))
    text = text.replace("{{STYLESHEET}}", html.escape(stylesheet.name, quote=True))
    text = text.replace("{{CONTENT}}", content)
    if "{{" in text or "<script" in text.lower():
        raise ValueError("WBS HTML contains unresolved or active content")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(text, encoding="utf-8")


def find_browser() -> str | None:
    candidates = [
        os.environ.get("WORK_WBS_BROWSER"),
        "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
        "/Applications/Chromium.app/Contents/MacOS/Chromium",
        shutil.which("google-chrome"), shutil.which("chromium"), shutil.which("chromium-browser"),
    ]
    return next((candidate for candidate in candidates if candidate and Path(candidate).is_file()), None)


def render_pdf(html_path: Path, output: Path) -> None:
    if output.exists():
        raise ValueError(f"refusing to overwrite existing PDF: {output}")
    browser = find_browser()
    if browser is None:
        raise RuntimeError("Chrome or Chromium was not found")
    with tempfile.TemporaryDirectory(prefix="work-wbs-chrome-") as profile:
        command = [browser, "--headless=new", "--disable-gpu", "--disable-extensions", "--allow-file-access-from-files", f"--user-data-dir={profile}", "--no-pdf-header-footer", f"--print-to-pdf={output}", html_path.resolve().as_uri()]
        with tempfile.TemporaryFile(mode="w+") as log:
            process = subprocess.Popen(command, stdout=log, stderr=subprocess.STDOUT, text=True, start_new_session=True)
            deadline = time.monotonic() + 30
            previous_size = -1; stable = 0
            while time.monotonic() < deadline:
                if output.is_file():
                    size = output.stat().st_size
                    stable = stable + 1 if size == previous_size else 0
                    previous_size = size
                    if size > 0 and stable >= 2:
                        break
                if process.poll() is not None:
                    break
                time.sleep(0.1)
            if process.poll() is None:
                os.killpg(process.pid, signal.SIGTERM)
                try: process.wait(timeout=3)
                except subprocess.TimeoutExpired:
                    os.killpg(process.pid, signal.SIGKILL); process.wait(timeout=3)
            log.seek(0); detail = log.read().strip()
    if not output.is_file() or not output.read_bytes().startswith(b"%PDF"):
        output.unlink(missing_ok=True)
        raise RuntimeError(f"Chrome PDF rendering failed: {detail or 'no PDF produced'}")


def verify_artifacts(manifest: dict[str, Any], xlsx: Path, html_path: Path, pdf: Path) -> None:
    expected_ids = {item["display_id"] for item in manifest["items"]}
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    sheet = workbook["WBS"]
    workbook_ids = {sheet.cell(row, 1).value for row in range(5, sheet.max_row + 1)}
    if workbook_ids != expected_ids:
        raise ValueError("XLSX work-item IDs do not match the manifest")
    text = html_path.read_text(encoding="utf-8")
    html_ids = set(re.findall(r'data-item-id="([^"]+)"', text))
    if html_ids != expected_ids:
        raise ValueError("HTML work-item IDs do not match the manifest")
    if not pdf.read_bytes().startswith(b"%PDF"):
        raise ValueError("PDF output is invalid")


def main() -> int:
    parser = argparse.ArgumentParser(description="Render a LUKUKU WBS manifest")
    parser.add_argument("manifest", type=Path)
    parser.add_argument("output_directory", type=Path)
    args = parser.parse_args()
    try:
        manifest = load_manifest(args.manifest)
        validate_manifest(manifest)
        output = args.output_directory.resolve()
        if output.exists() and any(output.iterdir()):
            raise ValueError(f"refusing to write into non-empty output directory: {output}")
        output.mkdir(parents=True, exist_ok=True)
        slug = _safe_slug(manifest["document"]["project_slug"])
        xlsx = output / f"{slug}-wbs.xlsx"
        pdf = output / f"{slug}-wbs.pdf"
        manifest_output = output / "wbs-manifest.json"
        manifest_output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        with tempfile.TemporaryDirectory(prefix="work-wbs-render-") as render_dir_value:
            render_dir = Path(render_dir_value)
            css = render_dir / "lukuku-wbs-template.css"
            shutil.copy2(SKILL_DIR / "assets" / "lukuku-wbs-template.css", css)
            html_path = render_dir / "wbs.html"
            render_workbook(manifest, xlsx)
            render_html(manifest, SKILL_DIR / "assets" / "lukuku-wbs-template.html", css, html_path)
            render_pdf(html_path, pdf)
            verify_artifacts(manifest, xlsx, html_path, pdf)
    except (OSError, RuntimeError, ValueError) as exc:
        print(str(exc), file=sys.stderr)
        return 1
    print(f"WBS artifacts created: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
